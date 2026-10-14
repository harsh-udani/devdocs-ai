
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


def extract_uploaded_document(
    filename: str,
    file_bytes: bytes,
) -> dict:
    """Extract text from a supported uploaded document."""

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. Upload PDF, DOCX, TXT, or MD."
        )

    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValueError("File size must be 10 MB or less.")

    if extension == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))

        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported.")

        text_parts = []

        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""

            if page_text.strip():
                text_parts.append(
                    f"[Page {page_number}]\n{page_text}"
                )

        text = "\n\n".join(text_parts)

    elif extension == ".docx":
        document = Document(BytesIO(file_bytes))

        parts = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Include table content in the extracted text.
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]
                parts.append(" | ".join(cells))

        text = "\n".join(parts)

    else:
        # Plain text and Markdown files.
        text = file_bytes.decode("utf-8-sig", errors="replace")

    if not text.strip():
        raise ValueError(
            "No readable text found. Scanned PDFs may need OCR."
        )

    return {
        "title": Path(filename).stem,
        "url": f"Uploaded file: {filename}",
        "text": text.strip(),
    }
